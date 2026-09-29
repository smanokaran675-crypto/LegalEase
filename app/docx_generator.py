from docx import Document
from docx.shared import Pt
import os


def create_docx(content):

    os.makedirs("generated", exist_ok=True)

    path = "generated/LegalEase_Document.docx"

    document = Document()

    title = document.add_paragraph()
    title.alignment = 1

    run = title.add_run("LegalEase Legal Document")
    run.bold = True
    run.font.size = Pt(16)

    document.add_paragraph()

    for line in content.splitlines():

        if line.strip():
            paragraph = document.add_paragraph()

            run = paragraph.add_run(line)
            run.font.size = Pt(11)

    document.save(path)

    return path